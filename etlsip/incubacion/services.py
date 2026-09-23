from __future__ import annotations

import os
import re
import uuid
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime, timedelta
from typing import Any
import calendar

from openpyxl import load_workbook
import pyodbc


SOURCE_CODES_DEFAULT = (
    "EOP B-HIM - Hatcher",
    "EOP B-HIM - Allocations",
)

HATCHERIES_DEFAULT = (
    "AVEGUAYAS",
    "AVEPICA",
    "INCA",
)


@dataclass(frozen=True)
class SqlServerConnectionConfig:
    server: str
    database: str
    username: str = ""
    password: str = ""
    trusted_connection: bool = True
    driver: str = "ODBC Driver 17 for SQL Server"
    timeout: int = 30


@dataclass(frozen=True)
class IncubacionETLConfig:
    source: SqlServerConnectionConfig
    destination: SqlServerConnectionConfig
    source_table: str
    destination_table: str
    destination_date_column: str = "xDate"
    source_codes: tuple[str, ...] = SOURCE_CODES_DEFAULT


@dataclass(frozen=True)
class CostoProdDetalleETLConfig:
    source: SqlServerConnectionConfig
    destination: SqlServerConnectionConfig
    source_table: str
    destination_table: str
    destination_date_column: str = "fecha_fin_mes"
    hatcheries: tuple[str, ...] = HATCHERIES_DEFAULT
    species_type: int = 1
    farm_type: int = 2


@dataclass(frozen=True)
class RecepcionETLConfig:
    source: SqlServerConnectionConfig
    destination: SqlServerConnectionConfig
    source_table: str
    destination_table: str
    destination_date_column: str = "Fecha_envio"
    egg_trans_code: int = 17
    facility_type: int = 1


@dataclass(frozen=True)
class OvoscopiaETLConfig:
    source: SqlServerConnectionConfig
    destination: SqlServerConnectionConfig
    destination_table: str = "dbo.Ovoscopia"


RECEPCION_NULL_COLUMNS = ("Cliente", "Tipo_documento", "No_documento")
RECEPCION_TIMESTAMP_COLUMN = "datemod"


@dataclass(frozen=True)
class PresupuestoETLConfig:
    destination: SqlServerConnectionConfig
    destination_table: str
    destination_date_column: str = "fecha_fin_mes"
    destination_key_column: str = "centro_costo"
    sheet_elemento_costo: str = "elemento_costo"
    sheet_presupuesto: str = "presupuesto"


@dataclass(frozen=True)
class ProteinJournalETLConfig:
    destination: SqlServerConnectionConfig
    destination_table: str
    destination_date_column: str = "xDate"
    destination_key_column: str = "centro_costo"
    destination_secondary_key_column: str = "granja_lote"
    sheet_data: str = "incubesa"
    sheet_mapping: str = "columnas"


@dataclass(frozen=True)
class CargasETLConfig:
    destination: SqlServerConnectionConfig
    destination_table: str = "dbo.cargas"
    sheet_data: str = "Hoja1"


@dataclass(frozen=True)
class VentaPollitoETLConfig:
    destination: SqlServerConnectionConfig
    destination_table: str = "dbo.VentaPollito"
    sheet_data: str = "Hoja1"


@dataclass(frozen=True)
class BajaPollitoETLConfig:
    destination: SqlServerConnectionConfig
    destination_table: str = "dbo.VentaPollito"
    sheet_data: str = "exportar"


_IDENTIFIER_PART_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def quote_identifier(identifier: str) -> str:
    parts = identifier.split(".")
    if not parts or any(not _IDENTIFIER_PART_PATTERN.match(part) for part in parts):
        raise ValueError(f"Invalid SQL identifier: {identifier}")
    return ".".join(f"[{part}]" for part in parts)


CARGAS_COLUMNS = (
    "No_incubadora", "No_granja", "No_lote", "No_maquina", "Fecha_carga",
    "Fecha_nacimiento", "No_carga", "Edad", "Semana", "Edad_huevo",
    "Huevos_cargados", "Huevos_transferidos", "Huevos_contaminados",
    "Pollos_nacidos", "Pollos_descarte", "Pollos_primera", "Pollos_eliminados",
    "Kg_HuevoCargado", "Kg_HuevoTransferido", "Kg_Pollos", "IRN", "Etapa",
    "Huevos_eliminados", "Fecha_transf", "Pollos_primera_H", "Pollos_primera_M",
    "Orden_prod", "Pollos_primera_Mixto", "Pollos_recuperados", "Huevo_recibido",
    "Categoria", "Tamano_huevo", "Tipo_maquina", "Fecha_postura", "No_sala",
)

CARGAS_KEY_COLUMNS = (
    "No_incubadora", "No_granja", "No_lote", "No_maquina", "Fecha_carga", "No_carga",
)

CARGAS_DATE_COLUMNS = {"Fecha_carga", "Fecha_nacimiento", "Fecha_transf", "Fecha_postura"}
CARGAS_INT_COLUMNS = {"No_maquina", "No_carga", "Edad_huevo"}
CARGAS_NUMERIC_COLUMNS = {
    "Edad", "Semana",
    "Huevos_cargados", "Huevos_transferidos", "Huevos_contaminados", "Pollos_nacidos",
    "Pollos_descarte", "Pollos_primera", "Pollos_eliminados", "Kg_HuevoCargado",
    "Kg_HuevoTransferido", "Kg_Pollos", "Huevos_eliminados", "Pollos_primera_H",
    "Pollos_primera_M", "Orden_prod", "Pollos_primera_Mixto", "Pollos_recuperados",
    "Huevo_recibido", "No_sala",
}

VENTA_POLLITO_COLUMNS = (
    "irn", "fecha", "cantidad", "fecha_semana", "tipo", "No_incubadora",
    "fecha_carga", "fecha_nacimiento", "semana_nacimiento", "granja_lote",
    "Cliente", "Tipo_documento", "No_documento", "no_ref",
)

VENTA_POLLITO_SOURCE_HEADERS = (
    "IRN", "DeliveryDate", "Quantity", "fecha_semana", "tipo", "HatcheryNo",
    "SetDate", "HatchDate", "ComplexEntityNo", "semana_nacimiento",
    "ComplexOrderNo", "BillToCustomerNo", "ComplexOrderNo", "RefNo",
)

VENTA_POLLITO_KEY_COLUMNS = ("irn",)

VENTA_POLLITO_DATE_COLUMNS = {
    "fecha", "fecha_semana", "fecha_carga", "fecha_nacimiento", "semana_nacimiento",
}

VENTA_POLLITO_INT_COLUMNS = {"cantidad", "no_ref"}

BAJA_POLLITO_SOURCE_COLUMNS = (
    "irn", "fecha", "cantidad", "fecha_semana", "tipo", "No_incubadora",
    "fecha_carga", "fecha_nacimiento", "semana_nacimiento", "granja_lote",
)


def build_connection_string(config: SqlServerConnectionConfig) -> str:
    base = [
        f"DRIVER={{{config.driver}}}",
        f"SERVER={config.server}",
        f"DATABASE={config.database}",
        f"Connection Timeout={config.timeout}",
    ]

    if config.trusted_connection:
        base.append("Trusted_Connection=yes")
    else:
        base.append(f"UID={config.username}")
        base.append(f"PWD={config.password}")

    return ";".join(base) + ";"


def open_connection(config: SqlServerConnectionConfig) -> pyodbc.Connection:
    connection_string = build_connection_string(config)
    return pyodbc.connect(connection_string)


def month_start(value: date) -> date:
    return date(value.year, value.month, 1)


def month_end(value: date) -> date:
    return date(value.year, value.month, calendar.monthrange(value.year, value.month)[1])


def build_extract_query(source_table: str, source_code_count: int) -> str:
    source_table_safe = quote_identifier(source_table)
    source_code_placeholders = ", ".join("?" for _ in range(source_code_count))

    return f"""
        SELECT
            xDate,
            EOMONTH(xDate) AS EndOfMonth,
            CostCenterNo AS centro_costo,
            ComplexEntityNo AS granja_lote,
            SystemCostElementNo AS elemento_costo,
            SystemLocationGroupNo AS grupo_ubicacion,
            RelativeAmount AS cantidad_relativa,
            SystemCostObjectNo AS objeto_costo,
            SourceCode AS codigo_fuente,
            SystemStageNo AS etapa_sistema,
            RelativeUnits AS unidades_relativas
        FROM {source_table_safe}
        WHERE IRN IS NOT NULL
          AND xDate BETWEEN ? AND ?
          AND SystemLocationGroupNo LIKE 'HAT %'
          AND SourceCode IN ({source_code_placeholders})
          AND SystemStageNo LIKE 'CHXPLT%'
    """


def build_costo_prod_detalle_extract_query(source_table: str, hatchery_count: int) -> str:
    if hatchery_count <= 0:
        raise ValueError("hatchery_count must be greater than zero")

    source_table_safe = quote_identifier(source_table)
    hatchery_placeholders = ", ".join("?" for _ in range(hatchery_count))

    return f"""
        SELECT
            be.GenerationCode AS codigo_generacion,
            r.CostCenterNo AS centro_costo,
            r.HatcheryNo AS no_incubadora,
            be.ComplexEntityNo AS granja_lote,
            (
                SELECT DivisionNo
                FROM mtech.ProteinDivisions
                WHERE IRN = pc.ProteinDivisionsIRN
            ) AS no_division,
            MAX(r.BegDate) AS fecha_inicio_mes,
            MAX(r.EndDate) AS fecha_fin_mes,
            SUM(r.ERStageAllocationAmount) AS er_stage_alloc,
            SUM(r.EROpsAllocationAmount) AS er_ops_alloc,
            SUM(r.EREggsPurchasedValue) AS er_compra_valor,
            SUM(r.EREggsPurchasedUnits) AS er_compra_unidades,
            SUM(r.EREggsReceivedValue) AS er_recepcion_valor,
            SUM(r.EREggsReceivedUnits) AS er_recepcion_unidades,
            SUM(r.EREggsDisposedUnits) AS er_eliminados,
            SUM(r.EREggsDowngradeUnits) AS er_bajas_calidad,
            SUM(r.EREggsTransferredInUnits) AS er_transfer_entrada,
            SUM(r.EREggsTransferredOutUnits) AS er_transfer_salida,
            SUM(r.EREggsSoldUnits) AS er_ventas_unidades,
            SUM(r.ERbeginningInventoryUnits) AS er_inv_inicial_unidades,
            SUM(r.ERbeginningInventoryValue) AS er_inv_inicial_valor,
            SUM(r.ERendingInventoryUnits) AS er_inv_final_unidades,
            SUM(r.ERendingInventoryValue) AS er_inv_final_valor,
            SUM(r.SRBeginningInventoryUnits) AS sr_inv_inicial_unidades,
            SUM(r.SRBeginningInventoryValue) AS sr_inv_inicial_valor,
            SUM(r.SREggsSetUnits) AS sr_hi_cargado_unidades,
            SUM(r.SREggsSetValue) AS sr_hi_cargado_valor,
            SUM(r.SREggsTransferredUnits) AS sr_transferencia_unidades,
            SUM(r.SREggsTransferredValue) AS sr_transferencia_valor,
            SUM(r.SREndingInventoryUnits) AS sr_inv_final_unidades,
            SUM(r.SREndingInventoryValue) AS sr_inv_final_valor,
            SUM(r.SROpsAllocationAmount) AS sr_distr_cif_valor,
            SUM(r.SRStageAllocationAmount) AS sr_distr_costos_etapa,
            SUM(r.SREggsDisposedUnits) AS sr_huevo_desecho_registrado_unidades,
            SUM(r.SRBeginningInventoryUnits)
                + SUM(r.SREggsSetUnits)
                - SUM(r.SREggsTransferredUnits)
                - SUM(r.SREndingInventoryUnits) AS sr_huevo_desecho_calculado_unidades,
            SUM(r.HRBeginningInventoryUnits) AS hr_inv_inicial_unidades,
            SUM(r.HRBeginningInventoryValue) AS hr_inv_inicial_valor,
            SUM(r.HREggsTransferredInUnits) AS hr_transf_ingreso_unidades,
            SUM(r.HREggsTransferredInValue) AS hr_transf_ingreso_valor,
            SUM(r.HREggsHatchedUnits) AS hr_pollos_nacidos_unidades,
            SUM(r.HREggsHatchedValue) AS hr_pollos_nacidos_valor,
            SUM(r.HREndingInventoryUnits) AS hr_inv_final_unidades,
            SUM(r.HREndingInventoryValue) AS hr_inv_final_valor,
            SUM(r.HRStageAllocationAmount) AS hr_costos_etapa_valor,
            SUM(r.HROpsAllocationAmount) AS hr_dist_cif_valor,
            r.ChickBeginningInventoryUnits AS pollitos_inv_inicial,
            r.ChickBeginningInventoryValue AS pollitos_inv_inicial_valor,
            SUM(r.ChicksHatchedUnits) AS pollitos_nacidos,
            SUM(r.ChicksHatchedAmount) AS pollitos_nacidos_valor,
            SUM(r.ChicksTransferredInUnits) AS pollitos_transfer_entrada,
            SUM(r.ChicksTransferredInValue) AS pollitos_transfer_entrada_valor,
            SUM(r.ChicksTransferredOutUnits) AS pollitos_transfer_salida,
            SUM(r.ChicksTransferredOutValue) AS pollitos_transfer_salida_valor,
            SUM(r.ChicksPurchasedUnits) AS pollitos_comprados,
            SUM(r.ChicksDisposedUnits) AS pollitos_eliminados,
            SUM(r.ChicksOpsAllocationAmount) AS pollitos_distr_cif_valor,
            SUM(r.ChicksPlacedUnits) AS pollitos_alojados,
            SUM(r.ChicksPlacedValue) AS pollitos_alojados_valor,
            SUM(r.ChicksSoldUnits) AS pollitos_vendidos,
            SUM(r.ChicksSoldValue) AS pollitos_vendidos_valor,
            r.ChickEndingInventoryUnits AS pollitos_inv_final_unidades,
            r.ChickEndingInventoryValue AS pollitos_inv_final_valor
        FROM {source_table_safe} r
        LEFT JOIN mtech.ProteinCostCenters pc ON pc.IRN = r.ProteinCostCentersIRN
        LEFT JOIN mtech.ProteinEntities pe ON pe.IRN = r.ProteinEntitiesIRN
        LEFT JOIN mtech.mvBimEntities be ON pe.IRN = be.ProteinEntitiesIRN
        WHERE r.HatcheryNo IN ({hatchery_placeholders})
          AND r.SpeciesType = ?
          AND r.FarmType = ?
          AND r.BegDate BETWEEN ? AND ?
          AND r.EndDate BETWEEN ? AND ?
        GROUP BY
            be.GenerationCode,
            be.ComplexEntityNo,
            r.CostCenterNo,
            r.HatcheryNo,
            pc.ProteinDivisionsIRN,
            r.ChickBeginningInventoryUnits,
            r.ChickBeginningInventoryValue,
            r.ChickEndingInventoryUnits,
            r.ChickEndingInventoryValue
    """


def extract_rows(
    source_conn: pyodbc.Connection,
    source_table: str,
    start_date: date,
    end_date: date,
    source_codes: tuple[str, ...],
) -> tuple[list[str], list[tuple[Any, ...]]]:
    query = build_extract_query(source_table=source_table, source_code_count=len(source_codes))
    params: list[Any] = [start_date, end_date, *source_codes]

    cursor = source_conn.cursor()
    cursor.execute(query, params)
    rows = cursor.fetchall()
    columns = [column[0] for column in cursor.description]
    cursor.close()

    return columns, [tuple(row) for row in rows]


def extract_rows_costo_prod_detalle(
    source_conn: pyodbc.Connection,
    source_table: str,
    start_date: date,
    end_date: date,
    hatcheries: tuple[str, ...],
    species_type: int,
    farm_type: int,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    query = build_costo_prod_detalle_extract_query(
        source_table=source_table,
        hatchery_count=len(hatcheries),
    )

    beg_start = month_start(start_date)
    beg_end = month_start(end_date)
    end_start = month_end(start_date)
    end_end = month_end(end_date)

    params: list[Any] = [
        *hatcheries,
        species_type,
        farm_type,
        beg_start,
        beg_end,
        end_start,
        end_end,
    ]

    cursor = source_conn.cursor()
    cursor.execute(query, params)
    rows = cursor.fetchall()
    columns = [column[0] for column in cursor.description]
    cursor.close()

    return columns, [tuple(row) for row in rows]


def build_recepcion_extract_query(source_table: str) -> str:
    source_table_safe = quote_identifier(source_table)

    return f"""
        SELECT
            HatcheryNo AS No_incubadora,
            FarmNo AS No_granja,
            ComplexEntityNo AS No_lote,
            TransDate AS Fecha_envio,
            Units AS Huevos_recibidos,
            DATEADD(day, 1 - DATEPART(weekday, TransDate), TransDate) AS Fecha_semana
        FROM {source_table_safe}
        WHERE EggTransCode = ?
          AND FacilityType = ?
          AND TransDate >= ?
          AND TransDate <= ?
    """


def extract_rows_recepcion(
    source_conn: pyodbc.Connection,
    source_table: str,
    start_date: date,
    end_date: date,
    egg_trans_code: int,
    facility_type: int,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    query = build_recepcion_extract_query(source_table=source_table)
    params: list[Any] = [egg_trans_code, facility_type, start_date, end_date]

    cursor = source_conn.cursor()
    cursor.execute(query, params)
    rows = cursor.fetchall()
    columns = [column[0] for column in cursor.description]
    cursor.close()

    return columns, [tuple(row) for row in rows]


def build_ovoscopia_extract_query() -> str:
    return """
        SELECT h.HatcheryNo AS No_incubadora,
               e.FarmNo AS No_granja,
               e.ComplexEntityNo AS No_lote,
               ht.SetDate AS Fecha_carga,
               ht.HatchDate AS Fecha_nacimiento,
               TRY_CONVERT(int, ht.LineSetCode) AS No_carga,
               CASE WHEN ht.BOType = 2 THEN 'Miraje' ELSE 'Residuos' END AS Tipo_registro,
               CEILING(DATEDIFF(day, e.AvgDatePlaced, ht.SetDate) / 7.0) AS Semana,
               ht.NumberOfSamples AS Huevos_muestra,
               ht.U_Infertil AS Huevos_infertiles,
               ht.U_Mort1 AS Mortalidad_1p,
               ht.U_Mort2 AS Mortalidad_2p,
               ht.U_Mort3 AS Mortalidad_3p,
               ht.U_Deforme AS Deformes,
               ht.U_Contaminado AS Contaminados,
               ht.U_PicadoNoNacido AS PicadoNoNacido,
               ht.IRN
        FROM mtech.HimBreakoutTrans ht
        INNER JOIN mtech.ProteinFacilityHatcheries h
            ON h.IRN = ht.ProteinFacilityHatcheriesIRN
        INNER JOIN mtech.mvBimEntities e
            ON e.IRN = ht.ProteinEntitiesIRN
        WHERE ht.HatchDate >= ? AND ht.HatchDate < ?
    """


def extract_rows_ovoscopia(
    source_conn: pyodbc.Connection,
    start_date: date,
    end_date: date,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    cursor = source_conn.cursor()
    end_exclusive = end_date + timedelta(days=1)
    cursor.execute(build_ovoscopia_extract_query(), start_date, end_exclusive)
    rows = cursor.fetchall()
    columns = [column[0] for column in cursor.description]
    cursor.close()
    return columns, [tuple(row) for row in rows]


def delete_destination_range(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    destination_date_column: str,
    start_date: date,
    end_date: date,
) -> int:
    destination_table_safe = quote_identifier(destination_table)
    destination_date_column_safe = quote_identifier(destination_date_column)

    delete_sql = f"""
        DELETE FROM {destination_table_safe}
        WHERE {destination_date_column_safe} BETWEEN ? AND ?
    """

    cursor = destination_conn.cursor()
    cursor.execute(delete_sql, (start_date, end_date))
    deleted_rows = cursor.rowcount if cursor.rowcount != -1 else 0
    cursor.close()
    return deleted_rows


def get_destination_columns(
    destination_conn: pyodbc.Connection,
    destination_table: str,
) -> list[str]:
    """Get list of columns in destination table."""
    schema, table = destination_table.split(".")
    
    cursor = destination_conn.cursor()
    query = f"""
    SELECT COLUMN_NAME
    FROM INFORMATION_SCHEMA.COLUMNS
    WHERE TABLE_SCHEMA = '{schema}' AND TABLE_NAME = '{table}'
    ORDER BY ORDINAL_POSITION
    """
    
    cursor.execute(query)
    columns = [row[0] for row in cursor.fetchall()]
    cursor.close()
    
    return columns


def get_destination_column_details(
    destination_conn: pyodbc.Connection,
    destination_table: str,
) -> dict[str, dict[str, Any]]:
    """Return destination column metadata keyed by lower-case column name."""
    schema, table = destination_table.split(".")

    cursor = destination_conn.cursor()
    query = """
    SELECT COLUMN_NAME, DATA_TYPE, NUMERIC_PRECISION, NUMERIC_SCALE
    FROM INFORMATION_SCHEMA.COLUMNS
    WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ?
    ORDER BY ORDINAL_POSITION
    """
    cursor.execute(query, (schema, table))
    rows = cursor.fetchall()
    cursor.close()

    details: dict[str, dict[str, Any]] = {}
    for column_name, data_type, precision, scale in rows:
        details[column_name.lower()] = {
            "name": column_name,
            "data_type": data_type,
            "precision": precision,
            "scale": scale,
        }
    return details


def normalize_decimal_rows_for_destination(
    mapped_columns: list[str],
    mapped_rows: list[tuple[Any, ...]],
    destination_column_details: dict[str, dict[str, Any]],
) -> list[tuple[Any, ...]]:
    """Normalize Decimal values to destination scale to avoid precision conversion errors."""
    decimal_positions: list[tuple[int, int]] = []

    for index, column_name in enumerate(mapped_columns):
        column_meta = destination_column_details.get(column_name.lower())
        if not column_meta:
            continue

        data_type = str(column_meta.get("data_type", "")).lower()
        scale = column_meta.get("scale")
        if data_type in {"decimal", "numeric"} and scale is not None:
            decimal_positions.append((index, int(scale)))

    if not decimal_positions:
        return mapped_rows

    normalized_rows: list[tuple[Any, ...]] = []
    for row in mapped_rows:
        mutable_row = list(row)
        for column_index, scale in decimal_positions:
            value = mutable_row[column_index]
            if not isinstance(value, Decimal):
                continue

            quantizer = Decimal(1).scaleb(-scale)
            mutable_row[column_index] = value.quantize(quantizer, rounding=ROUND_HALF_UP)

        normalized_rows.append(tuple(mutable_row))

    return normalized_rows


def map_columns_and_rows(
    source_columns: list[str],
    source_rows: list[tuple[Any, ...]],
    destination_columns: list[str],
) -> tuple[list[str], list[tuple[Any, ...]]]:
    """Map source columns to destination columns and filter rows accordingly.
    
    Handles:
    - Columns that exist in both source and destination
    - Special mappings (e.g., cantidad_relativa -> valor_relativo)
    - Skips identity columns (id)
    """
    # Create a map of source column name to index
    source_col_index = {col.lower(): idx for idx, col in enumerate(source_columns)}

    # Canonical map helps matching names with different underscore/case styles
    source_canonical_index = {
        re.sub(r"_", "", col.lower()): idx
        for idx, col in enumerate(source_columns)
    }
    
    # Define special mappings (destination -> source)
    special_mappings = {
        'valor_relativo': 'cantidad_relativa',
        'SourceCode': 'codigo_fuente',
        'SystemStageNo': 'etapa_sistema',
        'fecha_fin_mes': 'EndOfMonth',
        'grupo_costo': 'grupo_ubicacion',
        'SR_HuevoDesecho_Unidades': 'sr_huevo_desecho_calculado_unidades',
    }
    
    # Find which destination columns can be populated
    mapped_columns = []
    mapped_indices = []
    
    for dest_col in destination_columns:
        # Skip identity columns
        if dest_col.lower() == 'id':
            continue
        
        # Check for special mappings first
        if dest_col in special_mappings:
            source_col_name = special_mappings[dest_col]
            source_idx = source_col_index.get(source_col_name.lower())
            if source_idx is not None:
                mapped_columns.append(dest_col)
                mapped_indices.append(source_idx)
                continue
        
        # Check for direct column name match
        source_idx = source_col_index.get(dest_col.lower())
        if source_idx is not None:
            mapped_columns.append(dest_col)
            mapped_indices.append(source_idx)
            continue

        # Fallback to canonical matching (ignore underscores and case)
        canonical_dest = re.sub(r"_", "", dest_col.lower())
        source_idx = source_canonical_index.get(canonical_dest)
        if source_idx is not None:
            mapped_columns.append(dest_col)
            mapped_indices.append(source_idx)
    
    # Filter rows to only include mapped columns
    mapped_rows = []
    for row in source_rows:
        mapped_row = tuple(row[idx] if idx < len(row) else None for idx in mapped_indices)
        mapped_rows.append(mapped_row)
    
    return mapped_columns, mapped_rows


def insert_rows(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    columns: list[str],
    rows: list[tuple[Any, ...]],
    batch_size: int,
    fast_executemany: bool = True,
) -> int:
    if not rows:
        return 0

    # Get destination table columns
    dest_columns = get_destination_columns(destination_conn, destination_table)
    
    # Map and filter columns
    mapped_columns, mapped_rows = map_columns_and_rows(columns, rows, dest_columns)
    
    if not mapped_columns:
        raise ValueError(
            f"No common columns between source {columns} and destination {dest_columns}"
        )

    destination_column_details = get_destination_column_details(destination_conn, destination_table)
    mapped_rows = normalize_decimal_rows_for_destination(
        mapped_columns=mapped_columns,
        mapped_rows=mapped_rows,
        destination_column_details=destination_column_details,
    )

    destination_table_safe = quote_identifier(destination_table)
    columns_safe = ", ".join(quote_identifier(column) for column in mapped_columns)
    placeholders = ", ".join("?" for _ in mapped_columns)

    insert_sql = f"""
        INSERT INTO {destination_table_safe} ({columns_safe})
        VALUES ({placeholders})
    """

    cursor = destination_conn.cursor()
    cursor.fast_executemany = fast_executemany

    inserted_rows = 0
    for index in range(0, len(mapped_rows), batch_size):
        chunk = mapped_rows[index : index + batch_size]
        cursor.executemany(insert_sql, chunk)
        inserted_rows += len(chunk)

    cursor.close()
    return inserted_rows


def strip_string_values(rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    """Trim leading/trailing whitespace from string values (e.g. CHAR columns padded by SQL Server)."""
    return [
        tuple(value.strip() if isinstance(value, str) else value for value in row)
        for row in rows
    ]


def insert_rows_recepcion(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    columns: list[str],
    rows: list[tuple[Any, ...]],
    batch_size: int,
) -> int:
    """Insert Recepcion rows, filling Cliente/Tipo_documento/No_documento with NULL
    and datemod with the current system date. irn (identity) is left to SQL Server."""
    if not rows:
        return 0

    dest_columns = get_destination_columns(destination_conn, destination_table)
    mapped_columns, mapped_rows = map_columns_and_rows(columns, rows, dest_columns)

    if not mapped_columns:
        raise ValueError(
            f"No common columns between source {columns} and destination {dest_columns}"
        )

    mapped_rows = strip_string_values(mapped_rows)

    destination_column_details = get_destination_column_details(destination_conn, destination_table)
    mapped_rows = normalize_decimal_rows_for_destination(
        mapped_columns=mapped_columns,
        mapped_rows=mapped_rows,
        destination_column_details=destination_column_details,
    )

    dest_columns_lower = {column.lower() for column in dest_columns}
    mapped_lower = {column.lower() for column in mapped_columns}

    extra_null_columns = [
        column
        for column in RECEPCION_NULL_COLUMNS
        if column.lower() in dest_columns_lower and column.lower() not in mapped_lower
    ]
    include_timestamp = (
        RECEPCION_TIMESTAMP_COLUMN.lower() in dest_columns_lower
        and RECEPCION_TIMESTAMP_COLUMN.lower() not in mapped_lower
    )

    final_columns = list(mapped_columns) + extra_null_columns
    if include_timestamp:
        final_columns.append(RECEPCION_TIMESTAMP_COLUMN)

    load_timestamp = datetime.now()
    final_rows: list[tuple[Any, ...]] = []
    for row in mapped_rows:
        extended_row = list(row) + [None] * len(extra_null_columns)
        if include_timestamp:
            extended_row.append(load_timestamp)
        final_rows.append(tuple(extended_row))

    destination_table_safe = quote_identifier(destination_table)
    columns_safe = ", ".join(quote_identifier(column) for column in final_columns)
    placeholders = ", ".join("?" for _ in final_columns)

    insert_sql = f"""
        INSERT INTO {destination_table_safe} ({columns_safe})
        VALUES ({placeholders})
    """

    cursor = destination_conn.cursor()
    cursor.fast_executemany = True

    inserted_rows = 0
    for index in range(0, len(final_rows), batch_size):
        chunk = final_rows[index : index + batch_size]
        cursor.executemany(insert_sql, chunk)
        inserted_rows += len(chunk)

    cursor.close()
    return inserted_rows


def run_etl(
    config: IncubacionETLConfig,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, int]:
    source_conn = open_connection(config.source)
    destination_conn = open_connection(config.destination)

    try:
        columns, rows = extract_rows(
            source_conn=source_conn,
            source_table=config.source_table,
            start_date=start_date,
            end_date=end_date,
            source_codes=config.source_codes,
        )

        if dry_run:
            return {
                "source_rows": len(rows),
                "deleted_rows": 0,
                "inserted_rows": 0,
            }

        destination_conn.autocommit = False
        deleted_rows = delete_destination_range(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            destination_date_column=config.destination_date_column,
            start_date=start_date,
            end_date=end_date,
        )
        inserted_rows = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()

        return {
            "source_rows": len(rows),
            "deleted_rows": deleted_rows,
            "inserted_rows": inserted_rows,
        }
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        source_conn.close()
        destination_conn.close()


def run_etl_costo_prod_detalle(
    config: CostoProdDetalleETLConfig,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, int]:
    source_conn = open_connection(config.source)
    destination_conn = open_connection(config.destination)

    range_start = month_end(start_date)
    range_end = month_end(end_date)

    try:
        columns, rows = extract_rows_costo_prod_detalle(
            source_conn=source_conn,
            source_table=config.source_table,
            start_date=start_date,
            end_date=end_date,
            hatcheries=config.hatcheries,
            species_type=config.species_type,
            farm_type=config.farm_type,
        )

        if dry_run:
            return {
                "source_rows": len(rows),
                "deleted_rows": 0,
                "inserted_rows": 0,
            }

        destination_conn.autocommit = False
        deleted_rows = delete_destination_range(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            destination_date_column=config.destination_date_column,
            start_date=range_start,
            end_date=range_end,
        )
        inserted_rows = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()

        return {
            "source_rows": len(rows),
            "deleted_rows": deleted_rows,
            "inserted_rows": inserted_rows,
        }
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        source_conn.close()
        destination_conn.close()


def run_etl_recepcion(
    config: RecepcionETLConfig,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, int]:
    source_conn = open_connection(config.source)
    destination_conn = open_connection(config.destination)

    try:
        columns, rows = extract_rows_recepcion(
            source_conn=source_conn,
            source_table=config.source_table,
            start_date=start_date,
            end_date=end_date,
            egg_trans_code=config.egg_trans_code,
            facility_type=config.facility_type,
        )

        if dry_run:
            return {
                "source_rows": len(rows),
                "deleted_rows": 0,
                "inserted_rows": 0,
            }

        destination_conn.autocommit = False
        deleted_rows = delete_destination_range(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            destination_date_column=config.destination_date_column,
            start_date=start_date,
            end_date=end_date,
        )
        inserted_rows = insert_rows_recepcion(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()

        return {
            "source_rows": len(rows),
            "deleted_rows": deleted_rows,
            "inserted_rows": inserted_rows,
        }
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        source_conn.close()
        destination_conn.close()


def run_etl_ovoscopia(
    config: OvoscopiaETLConfig,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, int]:
    source_conn = open_connection(config.source)
    destination_conn = open_connection(config.destination)
    try:
        columns, rows = extract_rows_ovoscopia(source_conn, start_date, end_date)
        if dry_run:
            return {"source_rows": len(rows), "deleted_rows": 0, "inserted_rows": 0}

        destination_conn.autocommit = False
        deleted_rows = delete_destination_range(
            destination_conn, config.destination_table, "Fecha_nacimiento", start_date, end_date
        )
        inserted_rows = insert_rows(
            destination_conn, config.destination_table, columns, rows, batch_size,
            fast_executemany=False,
        )
        destination_conn.commit()
        return {"source_rows": len(rows), "deleted_rows": deleted_rows, "inserted_rows": inserted_rows}
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        source_conn.close()
        destination_conn.close()


def parse_excel_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            raise ValueError("Empty date value")
        try:
            return datetime.fromisoformat(text).date()
        except ValueError:
            return datetime.strptime(text, "%Y-%m-%d").date()

    raise ValueError(f"Unsupported date value type: {type(value).__name__}")


def parse_excel_int(value: Any, field_name: str) -> int:
    if value in (None, ""):
        raise ValueError(f"Missing required value for {field_name}")

    try:
        decimal_value = Decimal(str(value))
    except Exception as exc:
        raise ValueError(f"Invalid numeric value for {field_name}: {value}") from exc

    return int(decimal_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_cargas_datetime(value: Any, field_name: str, required: bool) -> datetime | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"Missing required value for {field_name}")
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError(f"Invalid datetime value for {field_name}: {value}") from exc
    raise ValueError(f"Invalid datetime value for {field_name}: {value}")


def parse_cargas_decimal(value: Any, field_name: str, required: bool) -> Decimal | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"Missing required value for {field_name}")
        return None
    try:
        decimal_value = Decimal(str(value))
    except Exception as exc:
        raise ValueError(f"Invalid numeric value for {field_name}: {value}") from exc
    if decimal_value != decimal_value.quantize(Decimal("1")):
        raise ValueError(f"Expected an integer numeric value for {field_name}: {value}")
    return decimal_value


def parse_cargas_float(value: Any, field_name: str, required: bool) -> float | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"Missing required value for {field_name}")
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid numeric value for {field_name}: {value}") from exc


def normalize_header_name(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def read_presupuesto_excel(
    excel_path: str,
    sheet_elemento_costo: str,
    sheet_presupuesto: str,
) -> tuple[dict[str, int], list[str], list[tuple[Any, ...]]]:
    if not os.path.exists(excel_path):
        raise ValueError(f"Excel file does not exist: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True)

    if sheet_elemento_costo not in workbook.sheetnames:
        raise ValueError(f"Missing required sheet: {sheet_elemento_costo}")
    if sheet_presupuesto not in workbook.sheetnames:
        raise ValueError(f"Missing required sheet: {sheet_presupuesto}")

    elemento_sheet = workbook[sheet_elemento_costo]
    presupuesto_sheet = workbook[sheet_presupuesto]

    elemento_map: dict[str, int] = {}
    for row in elemento_sheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue

        column_name = row[1]
        elemento_id = row[2]
        if column_name in (None, "") or elemento_id in (None, ""):
            continue

        key = normalize_header_name(column_name)
        parsed_id = int(elemento_id)
        existing = elemento_map.get(key)
        if existing is not None and existing != parsed_id:
            raise ValueError(
                f"Duplicate mapping for column '{column_name}' with different ids: "
                f"{existing} and {parsed_id}"
            )
        elemento_map[key] = parsed_id

    if not elemento_map:
        raise ValueError("No valid mappings found in elemento_costo sheet")

    rows_iter = list(presupuesto_sheet.iter_rows(values_only=True))
    header_index = -1
    header_row: list[Any] = []
    required_headers = {"fecha", "centro_costo"}
    for index, row in enumerate(rows_iter):
        if not row:
            continue
        normalized = {normalize_header_name(value) for value in row if value not in (None, "")}
        if required_headers.issubset(normalized):
            header_index = index
            header_row = list(row)
            break

    if header_index == -1:
        raise ValueError("Could not locate presupuesto header row with Fecha and centro_costo")

    data_rows: list[tuple[Any, ...]] = []
    for row in rows_iter[header_index + 1 :]:
        if row is None:
            continue
        if not any(value not in (None, "") for value in row):
            continue

        padded = list(row)
        if len(padded) < len(header_row):
            padded.extend([None] * (len(header_row) - len(padded)))
        data_rows.append(tuple(padded[: len(header_row)]))

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    return elemento_map, headers, data_rows


def get_presupuesto_date_range(
    excel_path: str,
    sheet_elemento_costo: str,
    sheet_presupuesto: str,
) -> tuple[date, date]:
    _elemento_map, headers, rows = read_presupuesto_excel(
        excel_path=excel_path,
        sheet_elemento_costo=sheet_elemento_costo,
        sheet_presupuesto=sheet_presupuesto,
    )

    header_index: dict[str, int] = {
        normalize_header_name(header): index
        for index, header in enumerate(headers)
        if header not in (None, "")
    }
    fecha_idx = header_index.get("fecha")
    if fecha_idx is None:
        raise ValueError("Presupuesto sheet must include Fecha")

    dates: list[date] = []
    for row in rows:
        if fecha_idx >= len(row):
            continue
        value = row[fecha_idx]
        if value in (None, ""):
            continue
        try:
            dates.append(parse_excel_date(value))
        except Exception:
            continue

    if not dates:
        raise ValueError("No valid Fecha values found in presupuesto sheet")

    return min(dates), max(dates)


def read_protein_journal_excel(
    excel_path: str,
    sheet_data: str,
    sheet_mapping: str,
) -> tuple[dict[str, dict[str, Any]], list[str], list[tuple[Any, ...]]]:
    if not os.path.exists(excel_path):
        raise ValueError(f"Excel file does not exist: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True)

    if sheet_data not in workbook.sheetnames:
        raise ValueError(f"Missing required sheet: {sheet_data}")
    if sheet_mapping not in workbook.sheetnames:
        raise ValueError(f"Missing required sheet: {sheet_mapping}")

    data_sheet = workbook[sheet_data]
    mapping_sheet = workbook[sheet_mapping]

    mapping_rows = list(mapping_sheet.iter_rows(values_only=True))
    if not mapping_rows:
        raise ValueError("Mapping sheet is empty")

    mapping_header = [str(value).strip() if value is not None else "" for value in mapping_rows[0]]
    mapping_index = {
        normalize_header_name(header): index
        for index, header in enumerate(mapping_header)
        if header
    }

    required_mapping_headers = {
        "columna",
        "grupo_costo",
        "objeto_costo",
        "sourcecode",
        "systemstageno",
    }
    if not required_mapping_headers.issubset(set(mapping_index.keys())):
        missing = sorted(required_mapping_headers - set(mapping_index.keys()))
        raise ValueError(f"Mapping sheet is missing required columns: {missing}")

    elemento_map: dict[str, dict[str, Any]] = {}
    for row in mapping_rows[1:]:
        if not row:
            continue

        col_idx = mapping_index["columna"]
        column_name = row[col_idx] if col_idx < len(row) else None
        if column_name in (None, ""):
            continue

        elemento_costo = str(column_name).strip()
        key = normalize_header_name(elemento_costo)
        if not key:
            continue

        grupo_idx = mapping_index["grupo_costo"]
        objeto_idx = mapping_index["objeto_costo"]
        source_idx = mapping_index["sourcecode"]
        stage_idx = mapping_index["systemstageno"]

        grupo_costo = row[grupo_idx] if grupo_idx < len(row) else None
        objeto_costo = row[objeto_idx] if objeto_idx < len(row) else None
        source_code = row[source_idx] if source_idx < len(row) else None
        system_stage = row[stage_idx] if stage_idx < len(row) else None

        if any(value in (None, "") for value in (grupo_costo, objeto_costo, source_code, system_stage)):
            raise ValueError(
                f"Incomplete mapping metadata for elemento '{elemento_costo}' in sheet {sheet_mapping}"
            )

        new_mapping = {
            "elemento_costo": elemento_costo,
            "grupo_costo": str(grupo_costo).strip(),
            "objeto_costo": str(objeto_costo).strip(),
            "SourceCode": str(source_code).strip(),
            "SystemStageNo": str(system_stage).strip(),
        }

        existing_mapping = elemento_map.get(key)
        if existing_mapping is not None and existing_mapping != new_mapping:
            raise ValueError(
                f"Duplicate mapping for elemento '{elemento_costo}' with conflicting metadata"
            )

        elemento_map[key] = new_mapping

    if not elemento_map:
        raise ValueError("No valid mappings found in mapping sheet")

    rows_iter = list(data_sheet.iter_rows(values_only=True))
    header_index = -1
    header_row: list[Any] = []
    required_data_headers = {"fecha", "centro_costo", "granja_lote", "huevo incubable"}
    for index, row in enumerate(rows_iter):
        if not row:
            continue
        normalized = {normalize_header_name(value) for value in row if value not in (None, "")}
        if required_data_headers.issubset(normalized):
            header_index = index
            header_row = list(row)
            break

    if header_index == -1:
        raise ValueError(
            "Could not locate incubesa header row with Fecha, centro_costo, granja_lote and Huevo Incubable"
        )

    data_rows: list[tuple[Any, ...]] = []
    for row in rows_iter[header_index + 1 :]:
        if row is None:
            continue
        if not any(value not in (None, "") for value in row):
            continue

        padded = list(row)
        if len(padded) < len(header_row):
            padded.extend([None] * (len(header_row) - len(padded)))
        data_rows.append(tuple(padded[: len(header_row)]))

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    return elemento_map, headers, data_rows


def get_protein_journal_date_range(
    excel_path: str,
    sheet_data: str,
    sheet_mapping: str,
) -> tuple[date, date]:
    _elemento_map, headers, rows = read_protein_journal_excel(
        excel_path=excel_path,
        sheet_data=sheet_data,
        sheet_mapping=sheet_mapping,
    )

    header_index: dict[str, int] = {
        normalize_header_name(header): index
        for index, header in enumerate(headers)
        if header not in (None, "")
    }
    fecha_idx = header_index.get("fecha")
    if fecha_idx is None:
        raise ValueError("Incubesa sheet must include Fecha")

    dates: list[date] = []
    for row in rows:
        if fecha_idx >= len(row):
            continue
        value = row[fecha_idx]
        if value in (None, ""):
            continue
        try:
            dates.append(parse_excel_date(value))
        except Exception:
            continue

    if not dates:
        raise ValueError("No valid Fecha values found in incubesa sheet")

    return min(dates), max(dates)


def transform_protein_journal_rows(
    elemento_map: dict[str, dict[str, Any]],
    data_headers: list[str],
    data_rows: list[tuple[Any, ...]],
    start_date: date,
    end_date: date,
    load_timestamp: datetime,
) -> tuple[list[str], list[tuple[Any, ...]], list[tuple[date, Any]], dict[str, Any]]:
    header_index: dict[str, int] = {
        normalize_header_name(header): index
        for index, header in enumerate(data_headers)
        if header not in (None, "")
    }

    fecha_idx = header_index.get("fecha")
    centro_idx = header_index.get("centro_costo")
    granja_idx = header_index.get("granja_lote")
    hi_idx = header_index.get(normalize_header_name("Huevo Incubable"))
    if fecha_idx is None or centro_idx is None or granja_idx is None or hi_idx is None:
        raise ValueError(
            "Incubesa sheet must include Fecha, centro_costo, granja_lote and Huevo Incubable"
        )

    generated_rows: list[tuple[Any, ...]] = []
    delete_keys: set[tuple[date, Any, Any]] = set()
    mapped_columns: set[str] = set()
    ignored_columns: set[str] = set()
    invalid_numeric_values = 0
    skipped_out_of_range = 0
    source_rows = 0

    base_header_positions = {
        fecha_idx,
        centro_idx,
        granja_idx,
        hi_idx,
    }

    for row in data_rows:
        try:
            xdate = parse_excel_date(row[fecha_idx])
        except Exception:
            continue

        if xdate < start_date or xdate > end_date:
            skipped_out_of_range += 1
            continue

        source_rows += 1
        centro_costo = row[centro_idx]
        if centro_costo in (None, ""):
            continue

        granja_lote = row[granja_idx] if granja_idx < len(row) else None

        try:
            unidades_relativas = parse_excel_int(
                row[hi_idx] if hi_idx < len(row) else None,
                "Huevo Incubable",
            )
        except ValueError as exc:
            raise ValueError(
                f"Invalid Huevo Incubable value for centro_costo={centro_costo}, fecha={xdate}: {exc}"
            ) from exc

        for index, header in enumerate(data_headers):
            if index in base_header_positions:
                continue
            if not header:
                continue

            normalized_header = normalize_header_name(header)
            elemento_meta = elemento_map.get(normalized_header)
            if elemento_meta is None:
                ignored_columns.add(header)
                continue

            value = row[index] if index < len(row) else None
            if value in (None, ""):
                continue

            try:
                valor_relativo = Decimal(str(value))
            except Exception:
                invalid_numeric_values += 1
                continue

            mapped_columns.add(header)
            delete_keys.add((xdate, centro_costo, granja_lote))
            generated_rows.append(
                (
                    xdate,
                    month_end(xdate),
                    centro_costo,
                    granja_lote,
                    elemento_meta["elemento_costo"],
                    elemento_meta["grupo_costo"],
                    valor_relativo,
                    unidades_relativas,
                    elemento_meta["objeto_costo"],
                    elemento_meta["SourceCode"],
                    elemento_meta["SystemStageNo"],
                    load_timestamp,
                )
            )

    metrics = {
        "source_rows": source_rows,
        "transformed_rows": len(generated_rows),
        "mapped_columns": sorted(mapped_columns),
        "ignored_columns": sorted(ignored_columns),
        "invalid_numeric_values": invalid_numeric_values,
        "skipped_out_of_range": skipped_out_of_range,
        "delete_key_count": len(delete_keys),
    }

    destination_columns = [
        "xDate",
        "fecha_fin_mes",
        "centro_costo",
        "granja_lote",
        "elemento_costo",
        "grupo_costo",
        "valor_relativo",
        "unidades_relativas",
        "objeto_costo",
        "SourceCode",
        "SystemStageNo",
        "fecha_carga",
    ]
    return destination_columns, generated_rows, sorted(delete_keys), metrics


def transform_presupuesto_rows(
    elemento_map: dict[str, int],
    presupuesto_headers: list[str],
    presupuesto_rows: list[tuple[Any, ...]],
    start_date: date,
    end_date: date,
    load_timestamp: datetime,
) -> tuple[list[str], list[tuple[Any, ...]], list[tuple[date, Any]], dict[str, Any]]:
    header_index: dict[str, int] = {
        normalize_header_name(header): index
        for index, header in enumerate(presupuesto_headers)
        if header not in (None, "")
    }

    fecha_idx = header_index.get("fecha")
    centro_idx = header_index.get("centro_costo")
    hi_idx = header_index.get(normalize_header_name("Huevo Incubable"))
    pollitos_idx = header_index.get(normalize_header_name("Pollito aprovechable"))
    if fecha_idx is None or centro_idx is None:
        raise ValueError("Presupuesto sheet must include Fecha and centro_costo")
    if hi_idx is None or pollitos_idx is None:
        raise ValueError(
            "Presupuesto sheet must include Huevo Incubable and Pollito aprovechable"
        )

    generated_rows: list[tuple[Any, ...]] = []
    delete_keys: set[tuple[date, Any]] = set()
    mapped_columns: set[str] = set()
    ignored_columns: set[str] = set()
    invalid_numeric_values = 0
    skipped_out_of_range = 0
    source_rows = 0

    for row in presupuesto_rows:
        try:
            fecha_value = parse_excel_date(row[fecha_idx])
        except Exception:
            continue

        if fecha_value < start_date or fecha_value > end_date:
            skipped_out_of_range += 1
            continue

        source_rows += 1
        centro_costo = row[centro_idx]
        if centro_costo in (None, ""):
            continue

        try:
            hi_cargados = parse_excel_int(
                row[hi_idx] if hi_idx < len(row) else None,
                "Huevo Incubable",
            )
            pollitos_nacidos = parse_excel_int(
                row[pollitos_idx] if pollitos_idx < len(row) else None,
                "Pollito aprovechable",
            )
        except ValueError as exc:
            raise ValueError(
                f"Invalid required values for centro_costo={centro_costo}, fecha={fecha_value}: {exc}"
            ) from exc

        for index, header in enumerate(presupuesto_headers):
            if index in (fecha_idx, centro_idx, hi_idx, pollitos_idx):
                continue
            if not header:
                continue

            normalized_header = normalize_header_name(header)
            elemento_costo = elemento_map.get(normalized_header)
            if elemento_costo is None:
                ignored_columns.add(header)
                continue

            value = row[index] if index < len(row) else None
            if value in (None, ""):
                continue

            try:
                valor_relativo = Decimal(str(value))
            except Exception:
                invalid_numeric_values += 1
                continue

            mapped_columns.add(header)
            delete_keys.add((fecha_value, centro_costo))
            generated_rows.append(
                (
                    fecha_value,
                    centro_costo,
                    str(elemento_costo),
                    valor_relativo,
                    hi_cargados,
                    pollitos_nacidos,
                    load_timestamp,
                )
            )

    metrics = {
        "source_rows": source_rows,
        "transformed_rows": len(generated_rows),
        "mapped_columns": sorted(mapped_columns),
        "ignored_columns": sorted(ignored_columns),
        "invalid_numeric_values": invalid_numeric_values,
        "skipped_out_of_range": skipped_out_of_range,
        "delete_key_count": len(delete_keys),
    }

    destination_columns = [
        "fecha_fin_mes",
        "centro_costo",
        "elemento_costo",
        "valor_relativo",
        "hi_cargados",
        "pollitos_nacidos",
        "fecha_carga",
    ]
    return destination_columns, generated_rows, sorted(delete_keys), metrics


def delete_destination_by_keys(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    destination_date_column: str,
    destination_key_column: str,
    keys: list[tuple[Any, ...]],
    destination_secondary_key_column: str | None = None,
) -> int:
    if not keys:
        return 0

    destination_table_safe = quote_identifier(destination_table)
    destination_date_column_safe = quote_identifier(destination_date_column)
    destination_key_column_safe = quote_identifier(destination_key_column)
    destination_secondary_key_column_safe = (
        quote_identifier(destination_secondary_key_column)
        if destination_secondary_key_column
        else None
    )

    if destination_secondary_key_column_safe:
        delete_sql = f"""
            DELETE FROM {destination_table_safe}
            WHERE {destination_date_column_safe} = ?
              AND {destination_key_column_safe} = ?
              AND (
                    ({destination_secondary_key_column_safe} = ?)
                 OR ({destination_secondary_key_column_safe} IS NULL AND ? IS NULL)
              )
        """
    else:
        delete_sql = f"""
            DELETE FROM {destination_table_safe}
            WHERE {destination_date_column_safe} = ?
              AND {destination_key_column_safe} = ?
        """

    cursor = destination_conn.cursor()
    deleted_rows = 0
    for key in keys:
        if destination_secondary_key_column_safe:
            if len(key) != 3:
                raise ValueError("Expected delete key tuple of (date, key, secondary_key)")
            cursor.execute(delete_sql, (key[0], key[1], key[2], key[2]))
        else:
            if len(key) != 2:
                raise ValueError("Expected delete key tuple of (date, key)")
            cursor.execute(delete_sql, key)
        if cursor.rowcount and cursor.rowcount > 0:
            deleted_rows += cursor.rowcount

    cursor.close()
    return deleted_rows


def run_etl_presupuesto(
    config: PresupuestoETLConfig,
    excel_path: str,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, Any]:
    destination_conn = open_connection(config.destination)
    load_timestamp = datetime.now()

    try:
        elemento_map, headers, presupuesto_rows = read_presupuesto_excel(
            excel_path=excel_path,
            sheet_elemento_costo=config.sheet_elemento_costo,
            sheet_presupuesto=config.sheet_presupuesto,
        )

        columns, rows, keys, metrics = transform_presupuesto_rows(
            elemento_map=elemento_map,
            presupuesto_headers=headers,
            presupuesto_rows=presupuesto_rows,
            start_date=start_date,
            end_date=end_date,
            load_timestamp=load_timestamp,
        )

        result: dict[str, Any] = {
            "source_rows": metrics["source_rows"],
            "deleted_rows": 0,
            "inserted_rows": 0,
            "summary": {
                **metrics,
                "excel_path": excel_path,
                "sheet_elemento_costo": config.sheet_elemento_costo,
                "sheet_presupuesto": config.sheet_presupuesto,
                "load_timestamp": load_timestamp.isoformat(),
                "mapped_elemento_count": len(elemento_map),
            },
        }

        if dry_run:
            return result

        destination_conn.autocommit = False
        deleted_rows = delete_destination_by_keys(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            destination_date_column=config.destination_date_column,
            destination_key_column=config.destination_key_column,
            keys=keys,
        )
        inserted_rows = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()

        result["deleted_rows"] = deleted_rows
        result["inserted_rows"] = inserted_rows
        return result
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        destination_conn.close()


def run_etl_protein_journal(
    config: ProteinJournalETLConfig,
    excel_path: str,
    start_date: date,
    end_date: date,
    dry_run: bool,
    batch_size: int,
) -> dict[str, Any]:
    destination_conn: pyodbc.Connection | None = None
    load_timestamp = datetime.now()

    try:
        elemento_map, headers, data_rows = read_protein_journal_excel(
            excel_path=excel_path,
            sheet_data=config.sheet_data,
            sheet_mapping=config.sheet_mapping,
        )

        columns, rows, keys, metrics = transform_protein_journal_rows(
            elemento_map=elemento_map,
            data_headers=headers,
            data_rows=data_rows,
            start_date=start_date,
            end_date=end_date,
            load_timestamp=load_timestamp,
        )

        result: dict[str, Any] = {
            "source_rows": metrics["source_rows"],
            "deleted_rows": 0,
            "inserted_rows": 0,
            "summary": {
                **metrics,
                "excel_path": excel_path,
                "sheet_data": config.sheet_data,
                "sheet_mapping": config.sheet_mapping,
                "load_timestamp": load_timestamp.isoformat(),
                "mapped_elemento_count": len(elemento_map),
            },
        }

        if dry_run:
            return result

        destination_conn = open_connection(config.destination)
        destination_conn.autocommit = False
        deleted_rows = delete_destination_by_keys(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            destination_date_column=config.destination_date_column,
            destination_key_column=config.destination_key_column,
            destination_secondary_key_column=config.destination_secondary_key_column,
            keys=keys,
        )
        inserted_rows = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()

        result["deleted_rows"] = deleted_rows
        result["inserted_rows"] = inserted_rows
        return result
    except Exception:
        if destination_conn is not None:
            destination_conn.rollback()
        raise
    finally:
        if destination_conn is not None:
            destination_conn.close()


def read_venta_pollito_excel(
    excel_path: str,
    sheet_data: str,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    if not os.path.exists(excel_path):
        raise ValueError(f"Excel file does not exist: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        if sheet_data not in workbook.sheetnames:
            raise ValueError(f"Missing required sheet: {sheet_data}")
        rows = list(workbook[sheet_data].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("VentaPollito sheet is empty")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    expected = [normalize_header_name(header) for header in VENTA_POLLITO_SOURCE_HEADERS]
    actual = [normalize_header_name(header) for header in headers]
    if actual != expected:
        raise ValueError(
            "VentaPollito sheet has unexpected columns. "
            f"Expected: {', '.join(VENTA_POLLITO_SOURCE_HEADERS)}"
        )

    source_indexes = (0, 1, 2, 3, 4, 5, 6, 7, 9, 8, 11, None, 10, 13)
    return list(VENTA_POLLITO_COLUMNS), [
        tuple(
            None if index is None else (row[index] if index < len(row) else None)
            for index in source_indexes
        )
        for row in rows[1:]
    ]


def get_venta_pollito_date_range(excel_path: str, sheet_data: str) -> tuple[date, date]:
    columns, rows = read_venta_pollito_excel(excel_path, sheet_data)
    date_index = columns.index("fecha")
    dates = [parse_excel_date(row[date_index]) for row in rows]
    if not dates:
        raise ValueError("No valid fecha values found in VentaPollito sheet")
    return min(dates), max(dates)


def transform_venta_pollito_rows(
    source_columns: list[str],
    source_rows: list[tuple[Any, ...]],
) -> tuple[list[str], list[tuple[Any, ...]], list[tuple[Any, ...]], dict[str, Any]]:
    required_columns = {
        "irn", "fecha", "cantidad", "fecha_semana", "tipo", "No_incubadora",
        "fecha_carga", "fecha_nacimiento", "semana_nacimiento", "granja_lote",
        "Cliente", "No_documento", "no_ref",
    }
    rows: list[tuple[Any, ...]] = []
    keys: set[tuple[Any, ...]] = set()

    for row_number, source_row in enumerate(source_rows, start=2):
        values: list[Any] = []
        for column, value in zip(source_columns, source_row):
            if column == "Tipo_documento":
                parsed = None
            elif column in VENTA_POLLITO_DATE_COLUMNS:
                parsed = parse_excel_date(value) if value not in (None, "") else None
            elif column in VENTA_POLLITO_INT_COLUMNS:
                parsed = parse_excel_int(value, column) if value not in (None, "") else None
            else:
                if value in (None, "") and column in required_columns:
                    raise ValueError(f"Row {row_number}: missing required value for {column}")
                parsed = str(value).strip() if value not in (None, "") else None
            if parsed is None and column in required_columns:
                raise ValueError(f"Row {row_number}: missing required value for {column}")
            values.append(parsed)

        normalized_row = tuple(values)
        rows.append(normalized_row)
        keys.add((normalized_row[source_columns.index("irn")],))

    return list(source_columns), rows, sorted(keys, key=str), {
        "source_rows": len(source_rows),
        "transformed_rows": len(rows),
        "key_count": len(keys),
    }


def delete_venta_pollito_by_keys(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    keys: list[tuple[Any, ...]],
) -> int:
    if not keys:
        return 0
    table = quote_identifier(destination_table)
    delete_sql = f"DELETE FROM {table} WHERE {quote_identifier('irn')} = ?"
    cursor = destination_conn.cursor()
    deleted_rows = 0
    try:
        for key in keys:
            cursor.execute(delete_sql, key)
            if cursor.rowcount > 0:
                deleted_rows += cursor.rowcount
    finally:
        cursor.close()
    return deleted_rows


def run_etl_venta_pollito(
    config: VentaPollitoETLConfig,
    excel_path: str,
    dry_run: bool,
    batch_size: int,
) -> dict[str, Any]:
    source_columns, source_rows = read_venta_pollito_excel(excel_path, config.sheet_data)
    columns, rows, keys, metrics = transform_venta_pollito_rows(source_columns, source_rows)
    result = {
        "source_rows": metrics["source_rows"],
        "deleted_rows": 0,
        "inserted_rows": 0,
        "summary": {**metrics, "excel_path": excel_path, "sheet_data": config.sheet_data},
    }
    if dry_run:
        return result

    destination_conn = open_connection(config.destination)
    try:
        destination_conn.autocommit = False
        result["deleted_rows"] = delete_venta_pollito_by_keys(
            destination_conn, config.destination_table, keys
        )
        result["inserted_rows"] = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()
        return result
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        destination_conn.close()


def read_baja_pollito_excel(
    excel_path: str,
    sheet_data: str,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    if not os.path.exists(excel_path):
        raise ValueError(f"Excel file does not exist: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        if sheet_data not in workbook.sheetnames:
            raise ValueError(f"Missing required sheet: {sheet_data}")
        rows = list(workbook[sheet_data].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("Baja pollito sheet is empty")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    expected = [normalize_header_name(header) for header in BAJA_POLLITO_SOURCE_COLUMNS]
    actual = [normalize_header_name(header) for header in headers[:len(expected)]]
    if actual != expected:
        raise ValueError(
            "Baja pollito sheet has unexpected columns. "
            f"Expected: {', '.join(BAJA_POLLITO_SOURCE_COLUMNS)}"
        )

    data_rows = [row for row in rows[1:] if any(value not in (None, "") for value in row)]
    return list(BAJA_POLLITO_SOURCE_COLUMNS), [
        tuple(row[index] if index < len(row) else None for index in range(len(BAJA_POLLITO_SOURCE_COLUMNS)))
        for row in data_rows
    ]


def get_baja_pollito_date_range(excel_path: str, sheet_data: str) -> tuple[date, date]:
    columns, rows = read_baja_pollito_excel(excel_path, sheet_data)
    date_index = columns.index("fecha")
    dates = [parse_excel_date(row[date_index]) for row in rows]
    if not dates:
        raise ValueError("No valid fecha values found in Baja pollito sheet")
    return min(dates), max(dates)


def transform_baja_pollito_rows(
    source_columns: list[str],
    source_rows: list[tuple[Any, ...]],
) -> tuple[list[str], list[tuple[Any, ...]], list[tuple[Any, ...]], dict[str, Any]]:
    required_columns = set(BAJA_POLLITO_SOURCE_COLUMNS)
    rows: list[tuple[Any, ...]] = []
    keys: set[tuple[Any, ...]] = set()
    source_indexes = {column: index for index, column in enumerate(source_columns)}

    for row_number, source_row in enumerate(source_rows, start=2):
        values: list[Any] = []
        for column in BAJA_POLLITO_SOURCE_COLUMNS:
            value = source_row[source_indexes[column]]
            if column in VENTA_POLLITO_DATE_COLUMNS:
                parsed = parse_excel_date(value) if value not in (None, "") else None
            elif column in VENTA_POLLITO_INT_COLUMNS:
                parsed = parse_excel_int(value, column) if value not in (None, "") else None
            else:
                parsed = str(value).strip() if value not in (None, "") else None
            if parsed is None and column in required_columns:
                raise ValueError(f"Row {row_number}: missing required value for {column}")
            values.append(parsed)

        source_irn = values[source_indexes["irn"]]
        generated_irn = str(uuid.uuid5(uuid.NAMESPACE_URL, f"baja-pollito:{source_irn}"))
        lote = values[source_indexes["granja_lote"]]
        values.extend(("BAJA", "INTERNO", f"BAJA-{lote}", 0))
        normalized_row = (generated_irn, *values[1:])
        rows.append(normalized_row)
        keys.add((generated_irn,))

    return list(VENTA_POLLITO_COLUMNS), rows, sorted(keys, key=str), {
        "source_rows": len(source_rows),
        "transformed_rows": len(rows),
        "key_count": len(keys),
        "generated_fields": ["irn", "Cliente", "Tipo_documento", "No_documento", "no_ref"],
    }


def run_etl_baja_pollito(
    config: BajaPollitoETLConfig,
    excel_path: str,
    dry_run: bool,
    batch_size: int,
) -> dict[str, Any]:
    source_columns, source_rows = read_baja_pollito_excel(excel_path, config.sheet_data)
    columns, rows, keys, metrics = transform_baja_pollito_rows(source_columns, source_rows)
    result = {
        "source_rows": metrics["source_rows"],
        "deleted_rows": 0,
        "inserted_rows": 0,
        "summary": {**metrics, "excel_path": excel_path, "sheet_data": config.sheet_data},
    }
    if dry_run:
        return result

    destination_conn = open_connection(config.destination)
    try:
        destination_conn.autocommit = False
        result["deleted_rows"] = delete_venta_pollito_by_keys(
            destination_conn, config.destination_table, keys
        )
        result["inserted_rows"] = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()
        return result
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        destination_conn.close()


def calculate_data_summary(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    destination_date_column: str,
    start_date: date,
    end_date: date,
    batch_size: int,
) -> dict[str, Any]:
    """Calculate summary statistics of inserted data.
    
    Returns:
        Dictionary with:
        - total_rows: Total rows inserted
        - batches: List of batch summaries
        - overall_stats: Overall statistics by numeric column
    """
    cursor = destination_conn.cursor()
    
    try:
        # Get column information
        table_name_parts = destination_table.split('.')
        schema_name = table_name_parts[0] if len(table_name_parts) > 1 else 'dbo'
        table_name_only = table_name_parts[-1]
        
        # Get numeric columns
        column_query = f"""
        SELECT COLUMN_NAME, DATA_TYPE
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = '{schema_name}' 
            AND TABLE_NAME = '{table_name_only}'
            AND DATA_TYPE IN ('decimal', 'numeric', 'float', 'real', 'bigint', 'int', 'smallint')
        ORDER BY ORDINAL_POSITION
        """
        cursor.execute(column_query)
        numeric_columns = [row[0] for row in cursor.fetchall()]

        # Detect incubadora-like column if present in destination table.
        all_columns_query = f"""
        SELECT COLUMN_NAME
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = '{schema_name}'
            AND TABLE_NAME = '{table_name_only}'
        ORDER BY ORDINAL_POSITION
        """
        cursor.execute(all_columns_query)
        all_columns = [row[0] for row in cursor.fetchall()]

        incubadora_column = None
        incubadora_candidates = {"no_incubadora", "hatcheryno", "hatchery_no"}
        for col_name in all_columns:
            if col_name.lower() in incubadora_candidates:
                incubadora_column = col_name
                break
        
        if not numeric_columns:
            by_incubadora = []
            if incubadora_column:
                group_query = f"""
                SELECT {quote_identifier(incubadora_column)}, COUNT(*)
                FROM {quote_identifier(destination_table)}
                WHERE {quote_identifier(destination_date_column)} BETWEEN ? AND ?
                GROUP BY {quote_identifier(incubadora_column)}
                ORDER BY {quote_identifier(incubadora_column)}
                """
                cursor.execute(group_query, start_date, end_date)
                by_incubadora = [
                    {
                        "incubadora": row[0],
                        "row_count": int(row[1]) if row[1] else 0,
                    }
                    for row in cursor.fetchall()
                ]

            return {
                "total_rows": 0,
                "batches": [],
                "overall_stats": {},
                "by_incubadora": by_incubadora,
            }
        
        # Get total row count
        count_query = f"""
        SELECT COUNT(*)
        FROM {quote_identifier(destination_table)}
        WHERE {quote_identifier(destination_date_column)} 
            BETWEEN ? AND ?
        """
        cursor.execute(count_query, start_date, end_date)
        total_rows = cursor.fetchone()[0]
        
        # Calculate batch summaries
        batches = []
        num_batches = (total_rows + batch_size - 1) // batch_size
        
        for batch_num in range(num_batches):
            offset = batch_num * batch_size
            batch_end = min(offset + batch_size, total_rows)
            batch_row_count = batch_end - offset
            
            # Get stats for this batch
            batch_stats = {"batch_number": batch_num + 1, "row_count": batch_row_count}
            
            for col in numeric_columns:
                stats_query = f"""
                SELECT 
                    SUM(CAST({quote_identifier(col)} AS FLOAT)),
                    AVG(CAST({quote_identifier(col)} AS FLOAT)),
                    MIN({quote_identifier(col)}),
                    MAX({quote_identifier(col)})
                FROM (
                    SELECT {quote_identifier(col)},
                        ROW_NUMBER() OVER (ORDER BY {quote_identifier(destination_date_column)}) as row_num
                    FROM {quote_identifier(destination_table)}
                    WHERE {quote_identifier(destination_date_column)} BETWEEN ? AND ?
                ) t
                WHERE row_num > ? AND row_num <= ?
                """
                cursor.execute(stats_query, start_date, end_date, offset, batch_end)
                row = cursor.fetchone()
                if row and row[0] is not None:
                    batch_stats[col] = {
                        "sum": float(row[0]) if row[0] else 0,
                        "avg": float(row[1]) if row[1] else 0,
                        "min": float(row[2]) if isinstance(row[2], (int, float)) else 0,
                        "max": float(row[3]) if isinstance(row[3], (int, float)) else 0,
                    }
            
            batches.append(batch_stats)
        
        # Calculate overall statistics
        overall_stats = {}
        for col in numeric_columns:
            stats_query = f"""
            SELECT 
                SUM(CAST({quote_identifier(col)} AS FLOAT)),
                AVG(CAST({quote_identifier(col)} AS FLOAT)),
                MIN({quote_identifier(col)}),
                MAX({quote_identifier(col)}),
                COUNT(*)
            FROM {quote_identifier(destination_table)}
            WHERE {quote_identifier(destination_date_column)} BETWEEN ? AND ?
            """
            cursor.execute(stats_query, start_date, end_date)
            row = cursor.fetchone()
            if row and row[0] is not None:
                overall_stats[col] = {
                    "sum": float(row[0]) if row[0] else 0,
                    "avg": float(row[1]) if row[1] else 0,
                    "min": float(row[2]) if isinstance(row[2], (int, float)) else 0,
                    "max": float(row[3]) if isinstance(row[3], (int, float)) else 0,
                    "count": int(row[4]) if row[4] else 0,
                }
        
        by_incubadora = []
        if incubadora_column:
            group_query = f"""
            SELECT {quote_identifier(incubadora_column)}, COUNT(*)
            FROM {quote_identifier(destination_table)}
            WHERE {quote_identifier(destination_date_column)} BETWEEN ? AND ?
            GROUP BY {quote_identifier(incubadora_column)}
            ORDER BY {quote_identifier(incubadora_column)}
            """
            cursor.execute(group_query, start_date, end_date)
            by_incubadora = [
                {
                    "incubadora": row[0],
                    "row_count": int(row[1]) if row[1] else 0,
                }
                for row in cursor.fetchall()
            ]

        return {
            "total_rows": total_rows,
            "batches": batches,
            "overall_stats": overall_stats,
            "by_incubadora": by_incubadora,
        }
    
    finally:
        cursor.close()


def read_cargas_excel(
    excel_path: str,
    sheet_data: str,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    if not os.path.exists(excel_path):
        raise ValueError(f"Excel file does not exist: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        if sheet_data not in workbook.sheetnames:
            raise ValueError(f"Missing required sheet: {sheet_data}")
        worksheet = workbook[sheet_data]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("Cargas sheet is empty")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    header_index = {normalize_header_name(header): index for index, header in enumerate(headers) if header}
    missing = [column for column in CARGAS_COLUMNS if normalize_header_name(column) not in header_index]
    if missing:
        raise ValueError(f"Cargas sheet is missing columns: {', '.join(missing)}")

    ordered_rows = []
    for row in rows[1:]:
        ordered_rows.append(tuple(row[header_index[normalize_header_name(column)]] for column in CARGAS_COLUMNS))
    return list(CARGAS_COLUMNS), ordered_rows


def get_cargas_date_range(excel_path: str, sheet_data: str) -> tuple[date, date]:
    columns, rows = read_cargas_excel(excel_path, sheet_data)
    date_index = columns.index("Fecha_carga")
    dates = [
        parse_cargas_datetime(row[date_index], "Fecha_carga", required=True).date()
        for row in rows
    ]
    if not dates:
        raise ValueError("No valid Fecha_carga values found in cargas sheet")
    return min(dates), max(dates)


def transform_cargas_rows(
    source_columns: list[str],
    source_rows: list[tuple[Any, ...]],
) -> tuple[list[str], list[tuple[Any, ...]], list[tuple[Any, ...]], dict[str, Any]]:
    nullable_columns = {
        "Edad", "Semana", "Kg_HuevoCargado", "Kg_HuevoTransferido", "Kg_Pollos",
        "Etapa", "Huevos_eliminados", "Fecha_transf", "Pollos_primera_H",
        "Pollos_primera_M", "Orden_prod", "Pollos_primera_Mixto", "Pollos_recuperados",
        "Huevo_recibido", "Categoria", "Tamano_huevo", "Tipo_maquina", "Fecha_postura",
        "No_sala",
    }
    float_columns = {"Edad", "Semana", "Kg_HuevoCargado", "Kg_HuevoTransferido", "Kg_Pollos"}
    rows: list[tuple[Any, ...]] = []
    keys: set[tuple[Any, ...]] = set()

    for row_number, source_row in enumerate(source_rows, start=2):
        values = []
        for column, value in zip(source_columns, source_row):
            required = column not in nullable_columns
            if column in CARGAS_DATE_COLUMNS:
                parsed = parse_cargas_datetime(value, column, required)
            elif column in CARGAS_INT_COLUMNS:
                parsed = parse_excel_int(value, column) if required or value not in (None, "") else None
            elif column in CARGAS_NUMERIC_COLUMNS:
                if column in float_columns:
                    parsed = parse_cargas_float(value, column, required)
                else:
                    parsed = parse_cargas_decimal(value, column, required)
            else:
                if value in (None, ""):
                    if required:
                        raise ValueError(f"Row {row_number}: missing required value for {column}")
                    parsed = None
                else:
                    parsed = str(value).strip()
            values.append(parsed)

        normalized_row = tuple(values)
        rows.append(normalized_row)
        keys.add(tuple(normalized_row[source_columns.index(column)] for column in CARGAS_KEY_COLUMNS))

    return source_columns, rows, sorted(keys, key=str), {
        "source_rows": len(source_rows),
        "transformed_rows": len(rows),
        "key_count": len(keys),
    }


def delete_cargas_by_keys(
    destination_conn: pyodbc.Connection,
    destination_table: str,
    keys: list[tuple[Any, ...]],
) -> int:
    if not keys:
        return 0
    table = quote_identifier(destination_table)
    predicates = " AND ".join(f"{quote_identifier(column)} = ?" for column in CARGAS_KEY_COLUMNS)
    delete_sql = f"DELETE FROM {table} WHERE {predicates}"
    cursor = destination_conn.cursor()
    deleted_rows = 0
    try:
        for key in keys:
            cursor.execute(delete_sql, key)
            if cursor.rowcount > 0:
                deleted_rows += cursor.rowcount
    finally:
        cursor.close()
    return deleted_rows


def run_etl_cargas(
    config: CargasETLConfig,
    excel_path: str,
    dry_run: bool,
    batch_size: int,
) -> dict[str, Any]:
    source_columns, source_rows = read_cargas_excel(excel_path, config.sheet_data)
    columns, rows, keys, metrics = transform_cargas_rows(source_columns, source_rows)
    result = {
        "source_rows": metrics["source_rows"],
        "deleted_rows": 0,
        "inserted_rows": 0,
        "summary": {**metrics, "excel_path": excel_path, "sheet_data": config.sheet_data},
    }
    if dry_run:
        return result

    destination_conn = open_connection(config.destination)
    try:
        destination_conn.autocommit = False
        result["deleted_rows"] = delete_cargas_by_keys(destination_conn, config.destination_table, keys)
        result["inserted_rows"] = insert_rows(
            destination_conn=destination_conn,
            destination_table=config.destination_table,
            columns=columns,
            rows=rows,
            batch_size=batch_size,
        )
        destination_conn.commit()
        return result
    except Exception:
        destination_conn.rollback()
        raise
    finally:
        destination_conn.close()
